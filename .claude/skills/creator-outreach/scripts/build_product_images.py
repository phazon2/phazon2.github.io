#!/usr/bin/env python3
"""Render the Gumroad cover and thumbnail for the pricing calculator.

Figures on the artwork are READ OUT OF THE BUILT WORKBOOK, never typed. The
first hand-made cover quoted £90.70 profit and £38.90 of direct costs where
the file computes £90.66 and £38.80 — three of eight numbers wrong, on the
image a buyer decides from. The same class of error put a wrong January net in
a draft email for product #1.

Image format is not a free choice either. Gumroad rejected a 2560x1440 PNG
with alpha at 2x density ("Could not process that image"). What it accepts is
a baseline JPEG at exactly the nominal size, so device_scale_factor is pinned
to 1 and the type is forced to jpeg.

    ./build_product_images.py --workbook Job-...-GBP.xlsx --out-dir gumroad/
"""
import argparse
import json
import pathlib
import shutil
import sys

from openpyxl import load_workbook

import verify_quote_calculator as V

CHROME = "/opt/pw-browsers/chromium-1194/chrome-linux/chrome"

SCENARIO = {
    "Type of work — copy a name from the Rate Card": "Patio — slabs",
    "Measurement (in that row's unit)": 48,
    "Access  (1.00 normal)": 1.20,
    "Condition  (1.00 normal)": 1.15,
    "Unknowns  (1.00 you have seen all of it)": 1.00,
    "Chemicals and consumables": 18,
    "Waste disposal / tip fees": 12,
    "Travel — round trip distance": 16,
    "Cost per unit distance": 0.55,
}
WANT = ("Rate from the Rate Card", "Combined multiplier", "Hours on site",
        "Total direct costs", "PRICE TO QUOTE",
        "Profit after costs and your own time", "Margin",
        "Effective hourly rate on this job", "Verdict")


def figures(workbook, tmp):
    """Apply the cover scenario to a copy and read the results back out."""
    shutil.copy(workbook, tmp)
    wb = load_workbook(tmp)
    ws = wb["Quote Builder"]
    lbl = {str(r[0].value).strip(): r[1].coordinate
           for r in ws.iter_rows() if r[0].value}
    missing = [k for k in SCENARIO if k not in lbl]
    if missing:
        raise SystemExit("labels missing from the workbook: " + repr(missing))
    for k, v in SCENARIO.items():
        ws[lbl[k]] = v
    wb.save(tmp)

    wb2 = load_workbook(tmp)
    rc = wb2["Rate Card"]
    V._RATE_TABLE.clear()
    for r in range(5, rc.max_row + 1):
        n = rc.cell(row=r, column=1).value
        if n:
            V._RATE_TABLE[str(n).strip()] = [
                n, rc.cell(row=r, column=2).value,
                rc.cell(row=r, column=3).value,
                rc.cell(row=r, column=4).value, None]
    s = V.Sheet(wb2)
    out = {}
    for k in WANT:
        c = lbl[k]
        out[k] = s.value("Quote Builder", c[0], int(c[1:]))
    return out


def shoot(html, png, w, h):
    from playwright.sync_api import sync_playwright
    with sync_playwright() as pw:
        b = pw.chromium.launch(executable_path=CHROME)
        pg = b.new_page(viewport={"width": w, "height": h},
                        device_scale_factor=1)
        pg.goto(f"file://{html}")
        pg.wait_for_timeout(1500)
        pg.screenshot(path=str(png), type="jpeg", quality=92)
        b.close()


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--workbook", required=True)
    ap.add_argument("--out-dir", required=True)
    ap.add_argument("--symbol", default="£")
    args = ap.parse_args()

    out = pathlib.Path(args.out_dir)
    out.mkdir(parents=True, exist_ok=True)
    tmp = out / "_scenario.xlsx"
    f = figures(args.workbook, tmp)
    tmp.unlink(missing_ok=True)
    (out / "cover-figures.json").write_text(json.dumps(f, indent=1))
    for k, v in f.items():
        print(f"  {k:<40}{v if isinstance(v, str) else format(v, ',.4f')}")

    cur = args.symbol
    templates = pathlib.Path(__file__).resolve().parent / "images"
    for name, (w, h) in (("cover", (1280, 720)), ("thumb", (1200, 1200))):
        src = templates / f"{name}.html"
        if not src.exists():
            print(f"  (no template {src}, skipping {name})", file=sys.stderr)
            continue
        html = src.read_text()
        for key, val in f.items():
            token = "{{" + key + "}}"
            if isinstance(val, str):
                html = html.replace(token, val)
            elif key == "Margin":
                html = html.replace(token, f"{val * 100:,.1f}%")
            elif key in ("Combined multiplier",):
                html = html.replace(token, f"{val:,.3f}")
            elif key == "Hours on site":
                html = html.replace(token, f"{val:,.2f}")
            else:
                html = html.replace(token, f"{cur}{val:,.2f}")
        html = html.replace("{{SYMBOL}}", cur)
        dst = out / f"{name}.html"
        dst.write_text(html)
        img = out / f"{name}-{w}x{h}.jpg"
        shoot(dst, img, w, h)
        print(f"  -> {img}")


if __name__ == "__main__":
    main()
